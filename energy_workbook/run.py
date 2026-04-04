#!/usr/bin/env python3
"""
Geopolitical Risk & Energy Markets: Complete Investment Research Workbook
========================================================================
Main entry point. Fetches data, builds all 20 tabs, saves to .xlsx.

Usage:
    python run.py
    python run.py --output my_workbook.xlsx
    python run.py --skip-fred          # Skip FRED data (no API key needed)
    python run.py --tabs 1 2 3         # Build only specific tabs

Requirements:
    pip install -r requirements.txt
    Set environment variable FRED_API_KEY for macro data (free at fred.stlouisfed.org)
"""
import argparse
import sys
import time
import traceback
from datetime import datetime
from pathlib import Path

import openpyxl
from openpyxl.styles import Font

# Add project root to path
sys.path.insert(0, str(Path(__file__).parent))

from energy_workbook.config import TAB_NAMES, TAB_COLORS
from energy_workbook.data_fetcher import fetch_all_data


def main():
    parser = argparse.ArgumentParser(description="Build Energy Finance Research Workbook")
    parser.add_argument("--output", "-o", default="Energy_Finance_Workbook.xlsx",
                        help="Output filename (default: Energy_Finance_Workbook.xlsx)")
    parser.add_argument("--skip-fred", action="store_true",
                        help="Skip FRED data fetching (creates empty macro columns)")
    parser.add_argument("--tabs", nargs="*", type=int,
                        help="Build only specific tabs (e.g., --tabs 1 2 3)")
    args = parser.parse_args()

    output_path = Path(__file__).parent / args.output
    tabs_to_build = set(args.tabs) if args.tabs else set(range(1, 21))

    print("=" * 70)
    print("  GEOPOLITICAL RISK & ENERGY MARKETS WORKBOOK BUILDER")
    print(f"  {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
    print("=" * 70)

    # ── Step 1: Fetch Data ──────────────────────────────────────────────
    t0 = time.time()
    if args.skip_fred:
        import os
        os.environ.pop("FRED_API_KEY", None)
        from energy_workbook import config
        config.FRED_API_KEY = "__SKIP__"

    data = fetch_all_data()
    print(f"\nData fetch completed in {time.time() - t0:.1f}s")

    # ── Step 2: Create Workbook ─────────────────────────────────────────
    wb = openpyxl.Workbook()
    # Remove default sheet
    wb.remove(wb.active)

    # Results dict — accumulates outputs from each tab for downstream use
    results = {}

    # ── Step 3: Build Tabs ──────────────────────────────────────────────
    tab_builders = _get_tab_builders()

    for tab_num in sorted(tabs_to_build):
        if tab_num not in tab_builders:
            print(f"\n  WARNING: No builder for Tab {tab_num}")
            continue

        builder_func, module_name = tab_builders[tab_num]
        tab_name = TAB_NAMES.get(tab_num, f"Tab{tab_num}")
        print(f"\n{'─' * 50}")
        print(f"  Building Tab {tab_num}: {tab_name}")
        print(f"{'─' * 50}")

        t1 = time.time()
        try:
            results = builder_func(wb, data, results)
            print(f"  ✓ Tab {tab_num} completed in {time.time() - t1:.1f}s")
        except Exception as e:
            print(f"  ✗ Tab {tab_num} FAILED: {e}")
            traceback.print_exc()
            # Create placeholder sheet
            ws = wb.create_sheet(title=tab_name)
            ws["A1"] = f"Tab {tab_num} failed to build: {str(e)}"
            ws["A1"].font = Font(color="CC0000", bold=True)

    # ── Step 4: Save ────────────────────────────────────────────────────
    print(f"\n{'=' * 70}")
    print(f"  Saving workbook to: {output_path}")

    try:
        wb.save(str(output_path))
        file_size = output_path.stat().st_size / (1024 * 1024)
        print(f"  ✓ Saved successfully ({file_size:.1f} MB)")
    except PermissionError:
        alt_path = output_path.with_stem(output_path.stem + "_new")
        print(f"  File locked. Saving to: {alt_path}")
        wb.save(str(alt_path))

    total_time = time.time() - t0
    print(f"\n  Total time: {total_time:.1f}s")
    print(f"  Sheets created: {len(wb.sheetnames)}")
    print("=" * 70)


def _get_tab_builders():
    """Import and return all tab builder functions."""
    builders = {}

    # Tabs 1-5: Data, Regime, Profiles, Performance, Correlation
    try:
        from energy_workbook.tabs_data import (
            build_tab_01, build_tab_02, build_tab_03,
            build_tab_04, build_tab_05,
        )
        builders[1] = (build_tab_01, "tabs_data")
        builders[2] = (build_tab_02, "tabs_data")
        builders[3] = (build_tab_03, "tabs_data")
        builders[4] = (build_tab_04, "tabs_data")
        builders[5] = (build_tab_05, "tabs_data")
    except ImportError as e:
        print(f"  WARNING: Could not import tabs_data: {e}")

    # Tabs 6-10: Regression, Volatility, Portfolio, Valuation, VaR
    try:
        from energy_workbook.tabs_models import (
            build_tab_06, build_tab_07, build_tab_08,
            build_tab_09, build_tab_10,
        )
        builders[6] = (build_tab_06, "tabs_models")
        builders[7] = (build_tab_07, "tabs_models")
        builders[8] = (build_tab_08, "tabs_models")
        builders[9] = (build_tab_09, "tabs_models")
        builders[10] = (build_tab_10, "tabs_models")
    except ImportError as e:
        print(f"  WARNING: Could not import tabs_models: {e}")

    # Tabs 11-15: Monte Carlo, Fixed Income, Options, Scenarios, Tail Risk
    try:
        from energy_workbook.tabs_risk import (
            build_tab_11, build_tab_12, build_tab_13,
            build_tab_14, build_tab_15,
        )
        builders[11] = (build_tab_11, "tabs_risk")
        builders[12] = (build_tab_12, "tabs_risk")
        builders[13] = (build_tab_13, "tabs_risk")
        builders[14] = (build_tab_14, "tabs_risk")
        builders[15] = (build_tab_15, "tabs_risk")
    except ImportError as e:
        print(f"  WARNING: Could not import tabs_risk: {e}")

    # Tabs 16-20: Transmission, Pairs, Factors, Backtest, Dashboard
    try:
        from energy_workbook.tabs_strategy import (
            build_tab_16, build_tab_17, build_tab_18,
            build_tab_19, build_tab_20,
        )
        builders[16] = (build_tab_16, "tabs_strategy")
        builders[17] = (build_tab_17, "tabs_strategy")
        builders[18] = (build_tab_18, "tabs_strategy")
        builders[19] = (build_tab_19, "tabs_strategy")
        builders[20] = (build_tab_20, "tabs_strategy")
    except ImportError as e:
        print(f"  WARNING: Could not import tabs_strategy: {e}")

    return builders


if __name__ == "__main__":
    main()
