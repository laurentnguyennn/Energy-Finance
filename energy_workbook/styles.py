"""
Excel formatting utilities using openpyxl.
"""
from openpyxl.styles import (
    Font, PatternFill, Alignment, Border, Side, numbers
)
from openpyxl.utils import get_column_letter
from openpyxl.formatting.rule import ColorScaleRule, CellIsRule

# ── Fonts ───────────────────────────────────────────────────────────────────
FONT_HEADER = Font(bold=True, size=11)
FONT_TITLE = Font(bold=True, size=14)
FONT_SUBTITLE = Font(bold=True, size=12)
FONT_INPUT = Font(color="0000CC", size=10)
FONT_LINKED = Font(color="006100", size=10)
FONT_ERROR = Font(color="CC0000", size=10)
FONT_OK = Font(color="006100", size=10)
FONT_NORMAL = Font(size=10)
FONT_SMALL = Font(size=9)

# ── Fills ───────────────────────────────────────────────────────────────────
FILL_INPUT = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
FILL_HEADER = PatternFill(start_color="D9D9D9", end_color="D9D9D9", fill_type="solid")
FILL_LIGHT_GREEN = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
FILL_LIGHT_RED = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
FILL_LIGHT_BLUE = PatternFill(start_color="D6E4F0", end_color="D6E4F0", fill_type="solid")
FILL_WHITE = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")

# ── Borders ─────────────────────────────────────────────────────────────────
THIN_BORDER = Border(
    left=Side(style="thin"), right=Side(style="thin"),
    top=Side(style="thin"), bottom=Side(style="thin"),
)
BOTTOM_BORDER = Border(bottom=Side(style="thin"))

# ── Alignments ──────────────────────────────────────────────────────────────
ALIGN_CENTER = Alignment(horizontal="center", vertical="center")
ALIGN_RIGHT = Alignment(horizontal="right", vertical="center")
ALIGN_LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)

# ── Number Formats ──────────────────────────────────────────────────────────
FMT_PCT = '0.00%'
FMT_PCT1 = '0.0%'
FMT_NUM2 = '0.00'
FMT_NUM3 = '0.000'
FMT_NUM4 = '0.0000'
FMT_MONEY = '$#,##0.00'
FMT_MONEY_M = '$#,##0.0,,"M"'
FMT_MONEY_B = '$#,##0.0,,,"B"'
FMT_INT = '#,##0'
FMT_DATE = 'YYYY-MM-DD'


def style_header_row(ws, row, max_col, font=FONT_HEADER, fill=FILL_HEADER):
    """Apply header formatting to a row."""
    for col in range(1, max_col + 1):
        cell = ws.cell(row=row, column=col)
        cell.font = font
        cell.fill = fill
        cell.border = BOTTOM_BORDER
        cell.alignment = ALIGN_CENTER


def style_data_range(ws, start_row, end_row, start_col, end_col, fmt=None):
    """Apply formatting to a data range."""
    for row in range(start_row, end_row + 1):
        for col in range(start_col, end_col + 1):
            cell = ws.cell(row=row, column=col)
            cell.font = FONT_NORMAL
            cell.border = THIN_BORDER
            if fmt:
                cell.number_format = fmt


def write_table(ws, data, start_row, start_col, headers=None, header_fmt=True,
                num_fmt=None, col_widths=None):
    """Write a 2D list/array to worksheet with optional headers and formatting.

    Args:
        ws: worksheet
        data: list of lists (rows × cols)
        start_row: starting row (1-indexed)
        start_col: starting column (1-indexed)
        headers: optional list of column headers
        header_fmt: apply header formatting
        num_fmt: number format string or list of per-column formats
        col_widths: optional list of column widths

    Returns:
        (end_row, end_col) tuple
    """
    row = start_row
    if headers:
        for j, h in enumerate(headers):
            cell = ws.cell(row=row, column=start_col + j, value=h)
            if header_fmt:
                cell.font = FONT_HEADER
                cell.fill = FILL_HEADER
                cell.border = BOTTOM_BORDER
                cell.alignment = ALIGN_CENTER
        row += 1

    for i, data_row in enumerate(data):
        for j, val in enumerate(data_row):
            cell = ws.cell(row=row + i, column=start_col + j, value=val)
            cell.font = FONT_NORMAL
            cell.border = THIN_BORDER
            if num_fmt:
                if isinstance(num_fmt, list):
                    if j < len(num_fmt) and num_fmt[j]:
                        cell.number_format = num_fmt[j]
                else:
                    cell.number_format = num_fmt

    end_row = row + len(data) - 1 if data else row
    end_col = start_col + (len(headers) - 1 if headers else (len(data[0]) - 1 if data else 0))

    if col_widths:
        for j, w in enumerate(col_widths):
            ws.column_dimensions[get_column_letter(start_col + j)].width = w

    return end_row, end_col


def add_correlation_coloring(ws, start_row, end_row, start_col, end_col):
    """Add conditional formatting color scale for correlation matrices."""
    cell_range = f"{get_column_letter(start_col)}{start_row}:{get_column_letter(end_col)}{end_row}"
    ws.conditional_formatting.add(
        cell_range,
        ColorScaleRule(
            start_type="num", start_value=-1, start_color="FF0000",
            mid_type="num", mid_value=0, mid_color="FFFFFF",
            end_type="num", end_value=1, end_color="00B050",
        ),
    )


def auto_width(ws, min_width=10, max_width=25):
    """Auto-adjust column widths based on content."""
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            if cell.value:
                max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[col_letter].width = min(max(max_len + 2, min_width), max_width)


def set_tab_color(ws, color_hex):
    """Set worksheet tab color."""
    ws.sheet_properties.tabColor = color_hex
