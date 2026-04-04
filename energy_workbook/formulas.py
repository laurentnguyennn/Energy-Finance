"""
Excel formula generation helpers for cross-sheet references.

Provides utilities to build Excel formula strings that reference cells across
worksheets, enabling the 20-sheet dependency chain required by the workbook spec.
"""
from openpyxl.utils import get_column_letter


# ── Cross-Sheet Reference Builders ─────────────────────────────────────────

def sheet_ref(sheet_name: str, col_letter: str, row: int) -> str:
    """Build a cross-sheet cell reference: 'SheetName'!A1"""
    return f"'{sheet_name}'!{col_letter}{row}"


def sheet_range(sheet_name: str, col_letter: str, start_row: int, end_row: int) -> str:
    """Build a cross-sheet range reference: 'SheetName'!A1:A100"""
    return f"'{sheet_name}'!{col_letter}{start_row}:{col_letter}{end_row}"


def local_range(col_letter: str, start_row: int, end_row: int) -> str:
    """Build a local range reference: A1:A100"""
    return f"{col_letter}{start_row}:{col_letter}{end_row}"


# ── Statistical Formula Builders ──────────────────────────────────────────

def f_average(rng: str) -> str:
    return f"=AVERAGE({rng})"


def f_stdev(rng: str) -> str:
    return f"=STDEV.S({rng})"


def f_ann_mean(rng: str) -> str:
    """Annualized mean return = AVERAGE(range) * 252"""
    return f"=AVERAGE({rng})*252"


def f_ann_vol(rng: str) -> str:
    """Annualized volatility = STDEV.S(range) * SQRT(252)"""
    return f"=STDEV.S({rng})*SQRT(252)"


def f_skew(rng: str) -> str:
    return f"=SKEW({rng})"


def f_kurt(rng: str) -> str:
    return f"=KURT({rng})"


def f_max(rng: str) -> str:
    return f"=MAX({rng})"


def f_min(rng: str) -> str:
    return f"=MIN({rng})"


def f_count(rng: str) -> str:
    return f"=COUNT({rng})"


def f_correl(rng1: str, rng2: str) -> str:
    return f"=CORREL({rng1},{rng2})"


def f_covariance(rng1: str, rng2: str, annualize: bool = True) -> str:
    mult = "*252" if annualize else ""
    return f"=COVARIANCE.S({rng1},{rng2}){mult}"


def f_countif(rng: str, criteria) -> str:
    return f"=COUNTIF({rng},{criteria})"


def f_averageif(criteria_rng: str, criteria, avg_rng: str) -> str:
    return f"=AVERAGEIF({criteria_rng},{criteria},{avg_rng})"


def f_averageifs(avg_rng: str, *criteria_pairs) -> str:
    """AVERAGEIFS(avg_range, criteria_range1, criteria1, criteria_range2, criteria2, ...)"""
    args = ",".join(f"{cr},{cv}" for cr, cv in criteria_pairs)
    return f"=AVERAGEIFS({avg_rng},{args})"


# ── Sharpe / Sortino ──────────────────────────────────────────────────────

def f_sharpe(mean_cell: str, vol_cell: str, rf: float) -> str:
    """Sharpe = (AnnMean - Rf) / AnnVol"""
    return f"=({mean_cell}-{rf})/{vol_cell}"


def f_sortino_denominator(rng: str) -> str:
    """Downside vol = SQRT(SUMPRODUCT(IF(range<0,range,0)^2)/COUNT(range))*SQRT(252)
    Must be entered as array formula or use implicit intersection."""
    return f"=SQRT(SUMPRODUCT((IF({rng}<0,{rng},0))^2)/COUNT({rng}))*SQRT(252)"


# ── Log Return ────────────────────────────────────────────────────────────

def f_log_return(price_cell_t: str, price_cell_prev: str) -> str:
    """LN(P_t / P_{t-1}) with blank handling"""
    return (f'=IF(AND({price_cell_t}<>"",{price_cell_prev}<>""),'
            f'LN({price_cell_t}/{price_cell_prev}),"")')


# ── VaR Formulas ──────────────────────────────────────────────────────────

def f_parametric_var(rng: str, confidence: float = 0.95) -> str:
    """Parametric VaR: -AVERAGE(range) + z * STDEV.S(range)"""
    z = {0.95: 1.645, 0.99: 2.326}.get(confidence, 1.645)
    return f"=-AVERAGE({rng})+{z}*STDEV.S({rng})"


def f_historical_var(rng: str, confidence: float = 0.95) -> str:
    """Historical VaR: -PERCENTILE(range, 1-confidence)"""
    alpha = round(1 - confidence, 4)
    return f"=-PERCENTILE({rng},{alpha})"


def f_cvar(rng: str, var_cell: str) -> str:
    """CVaR (Expected Shortfall): average of returns below -VaR threshold.
    This is an array formula approximation."""
    return f'=-AVERAGEIF({rng},"<-"&{var_cell})'


def f_cornish_fisher_z(skew_cell: str, kurt_cell: str, confidence: float = 0.95) -> str:
    """Cornish-Fisher adjusted z-score.
    z_CF = z + (z^2-1)/6*S + (z^3-3z)/24*K - (2z^3-5z)/36*S^2
    """
    z = {0.95: -1.645, 0.99: -2.326}.get(confidence, -1.645)
    return (f"={z}+({z}^2-1)/6*{skew_cell}"
            f"+({z}^3-3*{z})/24*{kurt_cell}"
            f"-(2*{z}^3-5*{z})/36*{skew_cell}^2")


# ── EWMA Volatility ──────────────────────────────────────────────────────

def f_ewma_init(return_rng: str) -> str:
    """Initialize EWMA variance from population variance of first 60 returns.
    VARP computes population variance (ddof=0) per RiskMetrics standard."""
    return f"=VARP({return_rng})"


def f_ewma_step(prev_var_cell: str, prev_return_cell: str, lam: float = 0.94) -> str:
    """EWMA variance recursion: sigma^2_t = lambda * sigma^2_{t-1} + (1-lambda) * r^2_{t-1}"""
    return f"={lam}*{prev_var_cell}+(1-{lam})*{prev_return_cell}^2"


def f_ewma_annualized_vol(var_cell: str) -> str:
    """Annualized EWMA volatility: SQRT(variance) * SQRT(252)"""
    return f"=SQRT({var_cell})*SQRT(252)"


# ── DCF Formulas ──────────────────────────────────────────────────────────

def f_dcf_revenue(prev_rev_cell: str, growth_cell: str, oil_ratio_cell: str) -> str:
    """Revenue = PrevRev * (1+growth) * (OilPrice_t / OilPrice_{t-1})"""
    return f"={prev_rev_cell}*(1+{growth_cell})*{oil_ratio_cell}"


def f_dcf_fcf(ebitda_cell: str, tax_cell: str, da_cell: str,
              capex_cell: str, wc_cell: str) -> str:
    """FCF = EBITDA*(1-Tax) + DA*Tax - Capex - deltaWC"""
    return f"={ebitda_cell}*(1-{tax_cell})+{da_cell}*{tax_cell}-{capex_cell}-{wc_cell}"


def f_dcf_pv(fcf_cell: str, wacc_cell: str, year: int) -> str:
    """PV = FCF / (1+WACC)^year"""
    return f"={fcf_cell}/(1+{wacc_cell})^{year}"


def f_dcf_terminal(fcf_cell: str, growth_cell: str, wacc_cell: str) -> str:
    """Terminal Value = FCF*(1+g)/(WACC-g)"""
    return f"={fcf_cell}*(1+{growth_cell})/({wacc_cell}-{growth_cell})"


def f_dcf_wacc(ev_cell: str, debt_cell: str, re_cell: str,
               rd_cell: str, tax_cell: str) -> str:
    """WACC = (E/V)*Re + (D/V)*Rd*(1-Tax), where V = E + D"""
    return (f"={ev_cell}/({ev_cell}+{debt_cell})*{re_cell}"
            f"+{debt_cell}/({ev_cell}+{debt_cell})*{rd_cell}*(1-{tax_cell})")


def f_capm(rf: float, beta_cell: str, erp: float) -> str:
    """Cost of equity: Re = Rf + Beta * ERP"""
    return f"={rf}+{beta_cell}*{erp}"


# ── Black-Scholes Formulas ────────────────────────────────────────────────

def f_bs_d1(s_cell: str, k_cell: str, t_cell: str,
            r_cell: str, sigma_cell: str, q_cell: str) -> str:
    """d1 = [LN(S/K) + (r-q+sigma^2/2)*T] / (sigma*SQRT(T))"""
    return (f"=(LN({s_cell}/{k_cell})+({r_cell}-{q_cell}+{sigma_cell}^2/2)*{t_cell})"
            f"/({sigma_cell}*SQRT({t_cell}))")


def f_bs_d2(d1_cell: str, sigma_cell: str, t_cell: str) -> str:
    """d2 = d1 - sigma * SQRT(T)"""
    return f"={d1_cell}-{sigma_cell}*SQRT({t_cell})"


def f_bs_call(s_cell: str, k_cell: str, t_cell: str,
              r_cell: str, q_cell: str, d1_cell: str, d2_cell: str) -> str:
    """Call = S*exp(-qT)*N(d1) - K*exp(-rT)*N(d2)"""
    return (f"={s_cell}*EXP(-{q_cell}*{t_cell})*NORM.S.DIST({d1_cell},TRUE)"
            f"-{k_cell}*EXP(-{r_cell}*{t_cell})*NORM.S.DIST({d2_cell},TRUE)")


def f_bs_put(s_cell: str, k_cell: str, t_cell: str,
             r_cell: str, q_cell: str, d1_cell: str, d2_cell: str) -> str:
    """Put = K*exp(-rT)*N(-d2) - S*exp(-qT)*N(-d1)"""
    return (f"={k_cell}*EXP(-{r_cell}*{t_cell})*NORM.S.DIST(-{d2_cell},TRUE)"
            f"-{s_cell}*EXP(-{q_cell}*{t_cell})*NORM.S.DIST(-{d1_cell},TRUE)")


def f_bs_delta_call(q_cell: str, t_cell: str, d1_cell: str) -> str:
    return f"=EXP(-{q_cell}*{t_cell})*NORM.S.DIST({d1_cell},TRUE)"


def f_bs_delta_put(q_cell: str, t_cell: str, d1_cell: str) -> str:
    return f"=EXP(-{q_cell}*{t_cell})*(NORM.S.DIST({d1_cell},TRUE)-1)"


def f_bs_gamma(s_cell: str, sigma_cell: str, t_cell: str,
               q_cell: str, d1_cell: str) -> str:
    return (f"=EXP(-{q_cell}*{t_cell})*NORM.S.DIST({d1_cell},FALSE)"
            f"/({s_cell}*{sigma_cell}*SQRT({t_cell}))")


def f_bs_vega(s_cell: str, t_cell: str, q_cell: str, d1_cell: str) -> str:
    """Vega per 1% vol change"""
    return f"={s_cell}*EXP(-{q_cell}*{t_cell})*NORM.S.DIST({d1_cell},FALSE)*SQRT({t_cell})/100"


# ── Bond Formulas ─────────────────────────────────────────────────────────

def f_bond_pv(coupon_cell: str, ytm_cell: str, face: float,
              n_periods: int) -> str:
    """Bond price = PV of coupons + PV of face.
    Uses Excel PV function for annuity + lump sum."""
    return (f"=-PV({ytm_cell}/2,{n_periods},{coupon_cell}*{face}/2,{face})")


def f_modified_duration(mac_dur_cell: str, ytm_cell: str) -> str:
    """Modified Duration = Macaulay Duration / (1 + y/2)"""
    return f"={mac_dur_cell}/(1+{ytm_cell}/2)"


def f_dv01(mod_dur_cell: str, price_cell: str) -> str:
    """DV01 = ModDur * Price * 0.0001"""
    return f"={mod_dur_cell}*{price_cell}*0.0001"


def f_price_sensitivity(mod_dur_cell: str, conv_cell: str,
                        price_cell: str, shift_bp: int) -> str:
    """Price change from rate shift using duration + convexity:
    NewPrice = Price + (-ModDur * dy * Price + 0.5 * Convexity * dy^2 * Price)"""
    dy = shift_bp / 10000
    return (f"={price_cell}+(-{mod_dur_cell}*{dy}*{price_cell}"
            f"+0.5*{conv_cell}*{dy}^2*{price_cell})")


# ── Pair Trade Formulas ───────────────────────────────────────────────────

def f_log_spread(price_a_cell: str, price_b_cell: str) -> str:
    """Log spread = LN(Price_A) - LN(Price_B)"""
    return f"=LN({price_a_cell})-LN({price_b_cell})"


def f_z_score(value_cell: str, mean_cell: str, std_cell: str) -> str:
    """Z-score = (Value - Mean) / StdDev"""
    return f"=({value_cell}-{mean_cell})/{std_cell}"


# ── Regime Classification Formula ─────────────────────────────────────────

def f_regime_classification(vix_cell: str, ovx_cell: str, s2s10_cell: str,
                            wti_cell: str, bev_cell: str, bws_cell: str,
                            ovx_90_val: str, wti_80_val: str) -> str:
    """Nested IF for regime classification (severity order: Crisis→Recession→Inflation→Growth).
    Simplified version — some conditions omitted for Excel formula length limits."""
    return (
        f'=IF(OR({ovx_cell}>{ovx_90_val},AND({vix_cell}>30,{bws_cell}>5)),3,'
        f'IF(AND({s2s10_cell}<0,{vix_cell}>25),4,'
        f'IF(AND({wti_cell}>{wti_80_val},{bev_cell}>2.5,{vix_cell}<30),2,'
        f'IF(AND({vix_cell}<20,{ovx_cell}<30,{s2s10_cell}>0),1,0))))'
    )


# ── Definition / Explanation Helper ───────────────────────────────────────

def add_definition(ws, row: int, col: int, title: str, definition: str,
                   expected_output: str = None, font=None, fill=None):
    """Write a definition/explanation block to the worksheet.

    Args:
        ws: worksheet
        row: starting row
        col: starting column
        title: metric or concept name
        definition: explanation text
        expected_output: what the user should expect to see
        font: optional font override
        fill: optional fill override

    Returns:
        next available row
    """
    from .styles import FONT_NORMAL, FILL_LIGHT_BLUE, ALIGN_LEFT
    ws.cell(row=row, column=col, value=f"DEF: {title}").font = font or FONT_NORMAL
    ws.cell(row=row, column=col).fill = fill or FILL_LIGHT_BLUE
    ws.cell(row=row, column=col).alignment = ALIGN_LEFT
    row += 1
    ws.cell(row=row, column=col, value=definition).font = FONT_NORMAL
    ws.cell(row=row, column=col).alignment = ALIGN_LEFT
    row += 1
    if expected_output:
        ws.cell(row=row, column=col, value=f"Expected: {expected_output}").font = FONT_NORMAL
        ws.cell(row=row, column=col).alignment = ALIGN_LEFT
        row += 1
    return row
